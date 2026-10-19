"""Keller (2014, PMID 24135711) sensitivity analysis + numerical fit diagnostics (v2).

INPUT  variants.csv with columns `variant` (CHROM_POS_MUTATION) and `exposure`
       (column name in the environmental file). One row = one variant x exposure
       pair; duplicates are dropped. If `exposure` is missing, cfg.exposure is
       used (with a warning).

GENETIC SOURCES: each generation is read from its own genetic file.
    gen 1   -> /mnt/cresla_prod/genome_datasets/merged_csv/gen.parquet
    gen 2,3 -> /mnt/cresla_prod/genome_datasets/gen3/gen2_gen3_selected_variants.parquet
(override with --source 1=PATH --source 23=PATH ...). Samples present in more than one
file are kept only in the generation they belong to (--no-dedup disables this). A guard
stops the run if two generations still share samples for the same exposure (--allow-overlap).

MODELS, on the SAME matched sample as the pipeline:
    nocov : onset ~ G*E
    base  : onset ~ G*E + C                    (pipeline model; C = sex + PCs)
    cxe   : onset ~ G*E + C + C:E              (intermediate step, reported only)
    full  : onset ~ G*E + C + C:E + C:G        (MAIN complete model)

DECISION RULES (all thresholds are in the THRESHOLDS block below, and saved in run_info.json):
    change  = |beta_full - beta_base| / |beta_base| >= CHANGE_MAIN  AND  bootstrap 95% CI of
              (beta_full - beta_base) excludes 0
    sign    = sign(beta_full) != sign(beta_base) AND bootstrap CI of the change excludes 0
    Sensitivity: the same rule at CHANGE_SENSITIVITY thresholds (table of counts).
    Validity: base and full fits well posed (rank, condition, finite HC3 SE) and SMD <= pipeline
              MAX_SMD. If not met the pair is NOT_EVALUABLE (no evidence against the variant).
    Significance with C:E / C:G is only reported, never used to exclude.
    Meta-analysis across cohorts on the FULL model: fixed and DerSimonian-Laird random effects,
    Cochran Q heterogeneity (p < HET_ALPHA flags the variant, never excludes it).

STATUS: STABLE | CHANGED | NO_CI | NOT_EVALUABLE | SKIP | FAIL (NO_CI = bootstrap CI missing) (FAIL = technical: paths disagree / error).

OUTPUT (default ./output/keller_sensitivity, cfg.keller_sensitivity_dir):
    results_all.csv, results_gen<N>.csv, summary_by_gen_exposure.csv, sensitivity_table.csv,
    meta_analysis.csv, all_variants_table.csv, keller_sensitivity_report.docx, figures/*.png,
    run_info.json

Usage:
    python keller_sensitivity.py variants.csv --generations 1 23
    python keller_sensitivity.py variants.csv --generations 1 23 --perm 500
    python keller_sensitivity.py variants.csv --boot 200      # quick test run
"""
from __future__ import annotations

import argparse
import contextlib
import dataclasses
import gc
import io
import json
import sys
import time
import warnings
import zlib
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import statsmodels.api as sm

# ============================================================================
# THRESHOLDS - single block. Conventional values, NOT tuned on our results.
# ============================================================================
THRESHOLDS = {
    # estimate change, full vs base
    "CHANGE_MAIN": 0.10,                 # classic confounder-change convention
    "CHANGE_SENSITIVITY": (0.20, 0.30),  # same rule, stricter thresholds
    # bootstrap on the matched sample (CI of the change full - base)
    "BOOT_B": 1000,
    "BOOT_SEED": 20240601,
    "BOOT_CI": 0.95,
    "BOOT_MIN_VALID_FRAC": 0.5,          # below this fraction of valid replicates the CI is not reported
    # cohort heterogeneity (Cochran Q): flagged, never a reason to exclude
    "HET_ALPHA": 0.05,
    # significance: reported only
    "SIG_ALPHA": 0.05,
    # numerical diagnostics (notes only, except well-posedness which defines validity)
    "COND_MAX": 1e6,
    "LEV_WARN": 0.5,
    "CONC_RTOL": 1e-6,
    "CONC_ATOL": 1e-8,
    # permutation health (--perm), notes only
    "PERM_VALID_WARN": 0.95,
    "PERM_VALID_FAIL": 0.50,
    "PERM_RANKDEF_WARN": 0.05,
    # SMD limit = cfg.max_smd of the pipeline (read at run time and stored in run_info.json)
}
T = THRESHOLDS
COND_MAX, LEV_WARN = T["COND_MAX"], T["LEV_WARN"]
CONC_RTOL, CONC_ATOL = T["CONC_RTOL"], T["CONC_ATOL"]
PERM_VALID_WARN, PERM_VALID_FAIL, PERM_RANKDEF_WARN = T["PERM_VALID_WARN"], T["PERM_VALID_FAIL"], T["PERM_RANKDEF_WARN"]

MODEL_TAGS = ("nocov", "base", "cxe", "full")
EVALUABLE = ("STABLE", "CHANGED", "NO_CI")
_BOOL_PREFIX = ("ok_", "same_sign_", "flag_")
_BOOL_NAMES = ("concordant", "boot_excl0", "evaluable")

GEN1_FILE = "/mnt/cresla_prod/genome_datasets/merged_csv/gen.parquet"
GEN23_FILE = "/mnt/cresla_prod/genome_datasets/gen3/gen2_gen3_selected_variants.parquet"
# generation -> genetic file. "23" is the gen2+gen3 pooled cohort (GENERATION in the .env).
DEFAULT_SOURCES = {1: GEN1_FILE, 23: GEN23_FILE, 2: GEN23_FILE, 3: GEN23_FILE}
PCA_BASE = "/mnt/cresla_prod/genome_datasets"
DEFAULT_PCA = {1: f"{PCA_BASE}/qc_output_gen1/pca_covariates.csv",
               23: f"{PCA_BASE}/qc_output_gen2_gen3/pca_covariates.csv",
               2: f"{PCA_BASE}/qc_output_gen2_gen3/pca_covariates.csv",
               3: f"{PCA_BASE}/qc_output_gen2_gen3/pca_covariates.csv"}
PCA_ATTR_CANDIDATES = ("pca_covariates_path_template", "PCA_COVARIATES_PATH_TEMPLATE", "pca_path_template")
RAW_ATTR_CANDIDATES = ("raw_file", "RAW_FILE", "raw_path", "genetic_file", "genetic_path")


# ============================================================================
# Fits and numerical diagnostics
# ============================================================================
def _design(df, variant_col, Ecols, Ccols, cxe: bool, gxc: bool):
    v = df[variant_col].to_numpy(float)
    E = df[Ecols].to_numpy(float)
    C = df[Ccols].to_numpy(float) if Ccols else np.empty((len(df), 0))
    cols, names = [np.ones(len(df))], ["Intercept"]
    cols.append(v); names.append("G")
    for j, e in enumerate(Ecols):
        cols.append(E[:, j]); names.append(e)
    for j, e in enumerate(Ecols):
        cols.append(v * E[:, j]); names.append(f"G:{e}")
    for j, c in enumerate(Ccols):
        cols.append(C[:, j]); names.append(c)
    if cxe:
        for j, c in enumerate(Ccols):
            for k, e in enumerate(Ecols):
                cols.append(C[:, j] * E[:, k]); names.append(f"{c}:{e}")
    if gxc:
        for j, c in enumerate(Ccols):
            cols.append(v * C[:, j]); names.append(f"G:{c}")
    return np.column_stack(cols), names


def _max_leverage(X: np.ndarray) -> float:
    """Maximum leverage via SVD (correct also for rank-deficient X)."""
    U, s, _ = np.linalg.svd(X, full_matrices=False)
    tol = s.max() * max(X.shape) * np.finfo(float).eps if s.size else 0.0
    keep = s > tol
    return float((U[:, keep] ** 2).sum(axis=1).max()) if keep.any() else float("nan")


def _fit(df, target, variant_col, Ecols, Ccols, cxe, gxc):
    y = df[target].to_numpy(float)
    X, names = _design(df, variant_col, Ecols, Ccols, cxe, gxc)
    Z = X[:, 1:]
    sd = Z.std(axis=0)
    Zs = (Z - Z.mean(axis=0)) / np.where(sd > 0, sd, 1)
    rank = int(np.linalg.matrix_rank(X))
    cond = float(np.linalg.cond(Zs)) if (sd > 0).all() else np.inf
    lev = _max_leverage(X)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        res = sm.OLS(y, X).fit(cov_type="HC3")
    i = names.index(f"G:{Ecols[0]}")
    se = float(res.bse[i])
    return {
        "beta": float(res.params[i]), "se": se, "p": float(res.pvalues[i]),
        "n_params": X.shape[1], "rank": rank, "cond": cond, "max_lev": lev,
        "ok": bool(rank == X.shape[1] and cond < COND_MAX and np.isfinite(se)),
    }


def fit_models(df, target, variant_col, Ecols, Ccols) -> dict:
    specs = {"nocov": ([], False, False), "base": (Ccols, False, False),
             "cxe": (Ccols, True, False), "full": (Ccols, True, True)}
    out = {}
    for tag, (cc, cxe, gxc) in specs.items():
        r = _fit(df, target, variant_col, Ecols, cc, cxe, gxc)
        out.update({f"{k}_{tag}": v for k, v in r.items()})
    b0, b1, b3 = out["beta_nocov"], out["beta_base"], out["beta_full"]
    out["delta_beta_cov"] = b1 - b0
    out["rel_shift_cov"] = (b1 - b0) / abs(b0) if b0 else np.nan
    out["same_sign_cov"] = bool(np.sign(b0) == np.sign(b1))
    out["delta_beta_full"] = b3 - b1
    out["rel_shift_full"] = (b3 - b1) / abs(b1) if b1 else np.nan
    out["same_sign_full"] = bool(np.sign(b1) == np.sign(b3))
    return out


def _bootstrap_change(m, target, col, Ecols, Ccols, B, seed) -> dict:
    """Bootstrap of the matched sample (rows resampled with replacement): CI of
    the change in the G:E estimate, full minus base."""
    y = m[target].to_numpy(float)
    Xb, nb = _design(m, col, Ecols, Ccols, False, False)
    Xf, _ = _design(m, col, Ecols, Ccols, True, True)
    i = nb.index(f"G:{Ecols[0]}")
    pb, pf = Xb.shape[1], Xf.shape[1]
    n = len(y)
    rng = np.random.RandomState(seed)
    d = []
    for _ in range(B):
        ix = rng.randint(0, n, n)
        bb, _, rb, _ = np.linalg.lstsq(Xb[ix], y[ix], rcond=None)
        if rb < pb:
            continue
        bf, _, rf, _ = np.linalg.lstsq(Xf[ix], y[ix], rcond=None)
        if rf < pf:
            continue
        d.append(bf[i] - bb[i])
    out = {"boot_B": B, "boot_valid": len(d), "boot_lo": np.nan, "boot_hi": np.nan,
           "boot_se": np.nan, "boot_excl0": False}
    if B and len(d) >= T["BOOT_MIN_VALID_FRAC"] * B and len(d) > 1:
        a = (1 - T["BOOT_CI"]) / 2 * 100
        lo, hi = np.percentile(d, [a, 100 - a])
        out.update({"boot_lo": float(lo), "boot_hi": float(hi), "boot_se": float(np.std(d, ddof=1)),
                    "boot_excl0": bool(lo > 0 or hi < 0)})
    return out


def _product_smd(df, treat_col, Ecols, Ccols) -> float:
    """Max SMD on the C x E PRODUCTS: matching balances main effects, not products."""
    worst, t = 0.0, df[treat_col] == 1
    for c in Ccols:
        for e in Ecols:
            p = df[c] * df[e]
            sd = np.sqrt((p[t].var() + p[~t].var()) / 2)
            if sd > 0:
                worst = max(worst, abs(p[t].mean() - p[~t].mean()) / sd)
    return float(worst)


def _carrier_support(m, treat_col, Ecols) -> dict:
    """How much information the carriers provide to identify G:E."""
    e = m[Ecols[0]]
    car = m[treat_col] == 1
    out = {
        "n_carriers_matched": int(car.sum()), "n_noncarriers_matched": int((~car).sum()),
        "n_distinct_E_carriers": int(e[car].nunique()),
        "sd_E_carriers": float(e[car].std()) if car.sum() > 1 else np.nan,
    }
    binary = e.nunique() <= 2
    out["E_binary"] = bool(binary)
    if binary:
        hi = e.max()
        out["n_carriers_E_high"] = int((car & (e == hi)).sum())
        out["n_carriers_E_low"] = int((car & (e != hi)).sum())
    return out


def _concordance(d, m, target, col, Ecols, Ccols, cfg) -> dict:
    """smf.ols (observed path) vs fast path (permutation path)."""
    import statsmodels.formula.api as smf
    from gene_environment.analysis.fast_ols import build_design_and_solve, interaction_column_index
    from gene_environment.analysis.matching import (
        match_control_units_indices, precompute_full_distance_matrix, precompute_scaled_covariates,
    )
    from gene_environment.analysis.modeling import _find_interaction_term, build_formula

    out = {"beta_smf": np.nan, "beta_fast": np.nan, "same_matched_set": False, "conc_error": ""}
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        mod = smf.ols(build_formula(target, col, Ecols, Ccols, m), data=m).fit()
    name = _find_interaction_term(mod.params.index, col)
    if name is not None:
        out["beta_smf"] = float(mod.params[name])

    try:
        Xs = precompute_scaled_covariates(d, Ecols + Ccols)
        D_full = precompute_full_distance_matrix(Xs)
        mi = match_control_units_indices(d["_match_variant"].to_numpy(), D_full, k=cfg.match_k)
        if mi is not None:
            idx = np.concatenate(mi)
            y = d[target].to_numpy(float)
            C = d[Ccols].to_numpy(float)[idx] if Ccols else None
            b = build_design_and_solve(d[col].to_numpy(float)[idx], d[Ecols].to_numpy(float)[idx], y[idx], C)
            if b is not None:
                out["beta_fast"] = float(b[interaction_column_index(len(Ecols))])
            out["same_matched_set"] = bool(
                len(idx) == len(m) and np.allclose(np.sort(y[idx]), np.sort(m[target].to_numpy(float)))
            )
    except Exception as exc:
        out["conc_error"] = f"{type(exc).__name__}: {exc}"
    return out


def _perm_health(d, obs_coef, col, Ecols, Ccols, cfg, B) -> dict:
    """First B permutations with the pipeline's seed and logic
    (modeling._run_permutation_batch), counting what goes wrong."""
    from gene_environment.analysis.fast_ols import build_design_and_solve, interaction_column_index
    from gene_environment.analysis.matching import (
        match_control_units_indices, precompute_full_distance_matrix, precompute_scaled_covariates,
    )
    from gene_environment.analysis.modeling import _stable_seed

    vv = d[col].to_numpy()
    y = d[cfg.target_col].to_numpy(float)
    E = d[Ecols].to_numpy(float)
    C = d[Ccols].to_numpy(float) if Ccols else None
    Xs = precompute_scaled_covariates(d, Ecols + Ccols)
    D_full = precompute_full_distance_matrix(Xs)
    inter = interaction_column_index(E.shape[1])
    q = 0 if C is None else C.shape[1]
    n_cols = 2 + 2 * E.shape[1] + q
    rng = np.random.RandomState(_stable_seed(cfg.random_state, col))

    cnt = {"matching": 0, "min_sample": 0, "design": 0}
    betas, rankdef = [], []
    for _ in range(B):
        pv = rng.permutation(vv)
        mi = match_control_units_indices((pv > 0).astype(int), D_full, k=cfg.match_k)
        if mi is None:
            cnt["matching"] += 1
            continue
        idx = np.concatenate(mi)
        if idx.shape[0] < cfg.min_sample_size:
            cnt["min_sample"] += 1
            continue
        Ci = C[idx] if C is not None else None
        beta = build_design_and_solve(pv[idx], E[idx], y[idx], Ci)
        if beta is None:
            cnt["design"] += 1
            continue
        v = pv[idx].astype(float).reshape(-1, 1)
        X = np.column_stack([np.ones(len(idx)), v, E[idx], v * E[idx]] + ([Ci] if Ci is not None else []))
        rankdef.append(int(np.linalg.matrix_rank(X)) < n_cols)
        betas.append(float(beta[inter]))

    n_valid = len(betas)
    out = {
        "perm_B": B, "perm_valid": n_valid,
        "perm_valid_frac": n_valid / B if B else np.nan,
        "perm_fail_matching": cnt["matching"], "perm_fail_min_sample": cnt["min_sample"],
        "perm_fail_design": cnt["design"],
    }
    if n_valid:
        b = np.array(betas)
        rd = np.array(rankdef)
        p = float(np.mean(np.abs(b) >= abs(obs_coef)))
        out.update({
            "perm_rankdef_frac": float(rd.mean()),
            "perm_p_emp": p,
            "perm_p_emp_fullrank": float(np.mean(np.abs(b[~rd]) >= abs(obs_coef))) if (~rd).any() else np.nan,
            "perm_mc_se": float(np.sqrt(max(p * (1 - p), 0) / n_valid)),
            "perm_mean": float(b.mean()), "perm_std": float(b.std()),
            "perm_z_obs": float((obs_coef - b.mean()) / b.std()) if b.std() > 0 else np.nan,
            "below_min_obs_coef": bool(abs(obs_coef) < cfg.min_obs_coef),
        })
    return out


# ============================================================================
# Classification
# ============================================================================
def classify(r: dict, has_perm: bool) -> tuple[str, str]:
    """Return (status, reasons). Validity first (NOT_EVALUABLE), then technical
    integrity (FAIL), then the change rules (CHANGED / STABLE)."""
    inval = []
    if not r.get("ok_base", False):
        inval.append("base fit not well posed")
    if not r.get("ok_full", False):
        inval.append("full fit not well posed")
    if r.get("max_smd_main", 0) > r.get("cfg_max_smd", np.inf):
        inval.append(f"SMD {r['max_smd_main']:.2f} > MAX_SMD {r['cfg_max_smd']:.2f}")
    if inval:
        return "NOT_EVALUABLE", "; ".join(inval)

    fail = []
    if r.get("conc_error"):
        fail.append(f"fast path could not be computed ({r['conc_error']})")
    elif r.get("concordant") is False:
        fail.append("smf.ols and fast path disagree")
    if not r.get("conc_error") and r.get("same_matched_set") is False:
        fail.append("observed matching != fast-path matching")
    if fail:
        return "FAIL", "; ".join(fail)

    notes = []
    if r.get("max_lev_full", 0) > LEV_WARN:
        notes.append(f"max leverage (full) {r['max_lev_full']:.2f} > {LEV_WARN}")
    if not r.get("ok_cxe", True):
        notes.append("C x E fit not well posed")
    if has_perm and "perm_valid_frac" in r:
        vf = r["perm_valid_frac"]
        if vf < PERM_VALID_FAIL:
            notes.append(f"only {vf:.0%} of permutations valid")
        elif vf < PERM_VALID_WARN:
            notes.append(f"{vf:.0%} of permutations valid")
        if (r.get("perm_rankdef_frac") or 0) > PERM_RANKDEF_WARN:
            notes.append(f"{r['perm_rankdef_frac']:.0%} of permutations rank-deficient")
    if r.get("flag_change"):
        notes.insert(0, f"estimate change {r['rel_shift_full']:+.0%} (CI excludes 0)")
    if r.get("flag_sign"):
        notes.insert(0, "sign inversion (CI excludes 0)")
    if not np.isfinite(r.get("boot_lo", np.nan)):
        # no bootstrap CI: the change rule cannot be applied, so never label the pair STABLE
        return "NO_CI", "; ".join(["bootstrap CI missing: change rule not assessable"] + notes)
    status = "CHANGED" if (r.get("flag_change") or r.get("flag_sign")) else "STABLE"
    return status, "; ".join(notes)


# ============================================================================
# One (variant, exposure) pair on one generation
# ============================================================================
def analyze_one(df, col, label, exposure, generation, Ecols, Ccols, cfg, perm_B, boot_B) -> dict:
    from gene_environment.analysis.matching import check_balance, match_control_units

    row = {"generation": generation, "exposure": exposure, "variant": label}
    d = df[df[col].notna() & (df[col] != ".")].copy()
    d[col] = pd.to_numeric(d[col]).astype(int)
    d["_match_variant"] = (d[col] > 0).astype(int)
    n_t, n_c = int(d["_match_variant"].sum()), int((d["_match_variant"] == 0).sum())
    row.update(n_carriers_all=n_t, n_noncarriers_all=n_c)
    # pipeline rule (not a threshold of this analysis)
    if n_t < cfg.min_treated or n_c < cfg.min_treated:
        return {**row, "status": "SKIP", "reasons": "below pipeline MIN_TREATED: the pipeline does not test it"}

    d = d[[cfg.target_col, col, "_match_variant"] + Ecols + Ccols].dropna()
    if d.shape[0] < cfg.min_sample_size:
        return {**row, "status": "SKIP", "reasons": "sample below MIN_SAMPLE_SIZE after dropna"}
    m = match_control_units(d, "_match_variant", k=cfg.match_k, covariates_for_matching=Ecols + Ccols)
    if m is None or m.shape[0] < cfg.min_sample_size:
        return {**row, "status": "SKIP", "reasons": "matching failed"}

    row["n_matched"] = int(m.shape[0])
    row.update(_carrier_support(m, "_match_variant", Ecols))
    row["max_smd_main"] = float(max(check_balance(m, "_match_variant", Ecols + Ccols).values() or [0]))
    row["cfg_max_smd"] = float(cfg.max_smd)
    row["max_smd_CxE"] = _product_smd(m, "_match_variant", Ecols, Ccols)
    row.update(fit_models(m, cfg.target_col, col, Ecols, Ccols))
    row["sig_base"] = bool(row["p_base"] < T["SIG_ALPHA"])
    row["sig_full"] = bool(row["p_full"] < T["SIG_ALPHA"])

    evaluable = bool(row["ok_base"] and row["ok_full"] and row["max_smd_main"] <= row["cfg_max_smd"])
    row["evaluable"] = evaluable
    if evaluable:
        seed = (T["BOOT_SEED"] + zlib.crc32(f"{label}|{exposure}|{generation}".encode())) % (2 ** 32)
        row.update(_bootstrap_change(m, cfg.target_col, col, Ecols, Ccols, boot_B, seed))
        rel = row["rel_shift_full"]
        row["flag_change"] = bool(np.isfinite(rel) and abs(rel) >= T["CHANGE_MAIN"] and row["boot_excl0"])
        row["flag_sign"] = bool((not row["same_sign_full"]) and row["boot_excl0"])

        conc = _concordance(d, m, cfg.target_col, col, Ecols, Ccols, cfg)
        row.update(conc)
        betas = [conc["beta_smf"], conc["beta_fast"], row["beta_base"]]
        if np.all(np.isfinite(betas)):
            ref = betas[0]
            row["conc_maxdiff"] = float(max(abs(b - ref) for b in betas))
            row["concordant"] = bool(row["conc_maxdiff"] <= CONC_ATOL + CONC_RTOL * abs(ref))
        else:
            row["conc_maxdiff"], row["concordant"] = np.nan, False
        if perm_B and np.isfinite(conc["beta_smf"]):
            row.update(_perm_health(d, conc["beta_smf"], col, Ecols, Ccols, cfg, perm_B))

    row["status"], row["reasons"] = classify(row, bool(perm_B))
    return row


# ============================================================================
# Data loading
# ============================================================================
def _quiet(fn, *a, **k):
    """Silence the pipeline's DEBUG/df.columns prints on stdout."""
    with contextlib.redirect_stdout(io.StringIO()):
        return fn(*a, **k)


def _read_input(path: str, cfg) -> pd.DataFrame:
    from gene_environment.logging_utils import get_logger
    log = get_logger(__name__)
    raw = pd.read_csv(path)
    if "variant" not in raw.columns:
        raise ValueError(f"{path}: missing 'variant' column (columns: {list(raw.columns)})")
    if "exposure" not in raw.columns:
        if not cfg.exposure:
            raise ValueError(f"{path}: missing 'exposure' column and EXPOSURE is not configured")
        log.warning("'exposure' column missing: using cfg.exposure=%s for all variants", cfg.exposure)
        raw["exposure"] = cfg.exposure
    raw["variant"] = raw["variant"].astype(str).str.strip()
    raw["exposure"] = raw["exposure"].astype(str).str.strip()
    raw = raw[(raw["variant"] != "") & (raw["exposure"] != "")]
    return raw[["variant", "exposure"]].drop_duplicates().reset_index(drop=True)


def _cfg_with_attr(cfg, names, value, what: str):
    for n in names:
        if hasattr(cfg, n):
            cur = getattr(cfg, n)
            return dataclasses.replace(cfg, **{n: Path(value) if isinstance(cur, Path) else value}), n
    raise AttributeError(f"cfg has no attribute for {what} among {list(names)}; "
                         f"cfg fields: {[f.name for f in dataclasses.fields(cfg)]}")


def _cfg_with_raw(cfg, path: str, raw_attr: str | None):
    """Copy of cfg pointing at another genetic file (attribute auto-detected or --raw-attr)."""
    names = [raw_attr] if raw_attr else list(RAW_ATTR_CANDIDATES)
    for n in names:
        if hasattr(cfg, n):
            cur = getattr(cfg, n)
            return dataclasses.replace(cfg, **{n: Path(path) if isinstance(cur, Path) else path}), n
    fields = [f.name for f in dataclasses.fields(cfg)]
    raise AttributeError(
        f"cfg has no attribute among {names}. Pass --raw-attr with the right one. "
        f"cfg fields containing 'raw' or 'file': {[f for f in fields if 'raw' in f.lower() or 'file' in f.lower()]}"
    )


def _pipeline_ids(cfg, path: str, raw_attr, loader) -> set:
    """Sample ids of a genetic file AS THE PIPELINE SEES THEM."""
    c2, _ = _cfg_with_raw(cfg, path, raw_attr)
    return set(loader(c2)[0]["id"].astype(str))


# ============================================================================
# Meta-analysis, summaries, sensitivity table
# ============================================================================
def _fitted(res: pd.DataFrame) -> pd.DataFrame:
    """Rows with real fit results (excludes SKIP and error rows without fit columns)."""
    if "ok_base" not in res.columns:
        return res.iloc[0:0].copy()
    fit = res[(res["status"] != "SKIP") & res["ok_base"].notna()].copy()
    for c in fit.columns:
        if c.startswith(_BOOL_PREFIX) or c in _BOOL_NAMES or c.startswith("sig_"):
            fit[c] = fit[c].fillna(False).astype(bool)
    return fit


def _evaluable(res: pd.DataFrame) -> pd.DataFrame:
    return _fitted(res[res["status"].isin(EVALUABLE)]) if len(res) else res.iloc[0:0].copy()


def meta_analysis(res: pd.DataFrame) -> pd.DataFrame:
    """Per (variant, exposure): meta-analysis of the FULL-model G:E estimates across the
    cohorts where the pair is evaluable. Fixed (inverse variance) and DerSimonian-Laird
    random effects; Cochran Q for heterogeneity (flagged, never an exclusion)."""
    from scipy import stats
    ev = _evaluable(res)
    rows = []
    if ev.empty:
        return pd.DataFrame(columns=["variant", "exposure", "k"])
    for (v, e), s in ev.groupby(["variant", "exposure"], sort=True):
        s = s[np.isfinite(s["se_full"]) & (s["se_full"] > 0)]
        k = len(s)
        if k == 0:
            continue
        b = s["beta_full"].to_numpy(float)
        w = 1.0 / s["se_full"].to_numpy(float) ** 2
        bf = float((w * b).sum() / w.sum())
        sef = float(np.sqrt(1.0 / w.sum()))
        r = {"variant": v, "exposure": e, "k": k,
             "generations": ",".join(map(str, s["generation"].tolist())),
             "beta_fixed": bf, "se_fixed": sef,
             "p_fixed": float(2 * stats.norm.sf(abs(bf / sef))),
             "Q": np.nan, "Q_df": k - 1, "Q_p": np.nan, "I2": np.nan, "tau2": np.nan,
             "beta_random": bf, "se_random": sef, "p_random": float(2 * stats.norm.sf(abs(bf / sef))),
             "flag_het": False}
        if k >= 2:
            q = float((w * (b - bf) ** 2).sum())
            r["Q"], r["Q_p"] = q, float(stats.chi2.sf(q, k - 1))
            r["I2"] = max(0.0, (q - (k - 1)) / q) if q > 0 else 0.0
            c = w.sum() - (w ** 2).sum() / w.sum()
            tau2 = max(0.0, (q - (k - 1)) / c) if c > 0 else 0.0
            wr = 1.0 / (1.0 / w + tau2)
            br = float((wr * b).sum() / wr.sum())
            ser = float(np.sqrt(1.0 / wr.sum()))
            r.update({"tau2": tau2, "beta_random": br, "se_random": ser,
                      "p_random": float(2 * stats.norm.sf(abs(br / ser))),
                      "flag_het": bool(r["Q_p"] < T["HET_ALPHA"])})
        rows.append(r)
    return pd.DataFrame(rows)


def sensitivity_table(res: pd.DataFrame, main_thr: float, sens_thrs) -> pd.DataFrame:
    """Counts of changed variants at each threshold, always with the CI condition
    (and, for reference, without it)."""
    ev = _evaluable(res)
    rows = []
    thrs = sorted({main_thr, *sens_thrs})
    for (g, e), s in ev.groupby(["generation", "exposure"], sort=True):
        rel = s["rel_shift_full"].abs()
        for thr in thrs:
            big = rel >= thr
            rows.append({"generation": g, "exposure": e, "threshold": thr, "is_main": thr == main_thr,
                         "n_evaluable": len(s),
                         "n_change_no_ci": int(big.sum()),
                         "n_change_ci": int((big & s["boot_excl0"]).sum()),
                         "n_sign_flip_ci": int(((~s["same_sign_full"]) & s["boot_excl0"]).sum())})
    return pd.DataFrame(rows)


def summarize(res: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for (g, e), s in res.groupby(["generation", "exposure"], sort=True):
        r = {"generation": g, "exposure": e,
             "n_cohort": (s["n_cohort"].dropna().iloc[0] if "n_cohort" in s and s["n_cohort"].notna().any() else np.nan),
             "n_pairs": len(s)}
        for st in ("STABLE", "CHANGED", "NO_CI", "NOT_EVALUABLE", "SKIP", "FAIL"):
            r[f"n_{st.lower()}"] = int((s["status"] == st).sum())
        ev = _evaluable(s)
        if len(ev):
            r["frac_concordant"] = float(ev["concordant"].mean())
            r["median_cond_full"] = float(ev["cond_full"].median())
            r["median_abs_shift_full"] = float(ev["rel_shift_full"].abs().median())
            r["median_se_ratio_full_base"] = float((ev["se_full"] / ev["se_base"]).median())
            r["n_sig_base"] = int(ev["sig_base"].sum())
            r["n_sig_full"] = int(ev["sig_full"].sum())
            r["n_sig_base_still_sig_full"] = int((ev["sig_base"] & ev["sig_full"]).sum())
        rows.append(r)
    return pd.DataFrame(rows)


_WIDE_COLS = ["status", "reasons", "n_matched", "n_carriers_matched", "max_smd_main",
              "beta_nocov", "beta_base", "se_base", "p_base", "beta_cxe", "beta_full", "se_full", "p_full",
              "delta_beta_full", "rel_shift_full", "boot_lo", "boot_hi", "boot_excl0",
              "flag_change", "flag_sign"]


def all_variants_table(res: pd.DataFrame, meta: pd.DataFrame) -> pd.DataFrame:
    """One row per variant x exposure, cohorts side by side, plus the meta-analysis."""
    base = res[["variant", "exposure"]].drop_duplicates().set_index(["variant", "exposure"])
    for g in sorted(res["generation"].unique()):
        s = res[res["generation"] == g].drop_duplicates(["variant", "exposure"]).set_index(["variant", "exposure"])
        base = base.join(s.reindex(columns=_WIDE_COLS).add_suffix(f"_g{g}"), how="left")
    if len(meta):
        base = base.join(meta.set_index(["variant", "exposure"]), how="left")
    return base.reset_index().sort_values(["exposure", "variant"]).reset_index(drop=True)


# ============================================================================
# Figures
# ============================================================================
STATUS_COLORS = {"STABLE": "#2a9d8f", "CHANGED": "#e76f51", "NO_CI": "#8d99ae", "NOT_EVALUABLE": "#e9c46a",
                 "FAIL": "#e63946", "SKIP": "#999999"}


def make_figures(res: pd.DataFrame, fig_dir: Path, has_perm: bool) -> dict[str, Path]:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig_dir.mkdir(parents=True, exist_ok=True)
    gens = sorted(res["generation"].unique())
    ev = _evaluable(res)
    fit = _fitted(res)
    out: dict[str, Path] = {}

    fig, ax = plt.subplots(figsize=(max(5, 1.3 * res.groupby(["generation", "exposure"]).ngroups), 4))
    tab = res.groupby(["generation", "exposure", "status"]).size().unstack(fill_value=0)
    tab = tab.reindex(columns=[c for c in STATUS_COLORS if c in tab.columns])
    tab.plot(kind="bar", stacked=True, ax=ax, color=[STATUS_COLORS[c] for c in tab.columns], width=0.75)
    ax.set_xlabel("(generation, exposure)"); ax.set_ylabel("no. of variants")
    ax.set_xticklabels([f"g{a}\n{b}" for a, b in tab.index], rotation=0, fontsize=8)
    ax.legend(title="status", frameon=False); fig.tight_layout()
    out["status"] = fig_dir / "status_counts.png"; fig.savefig(out["status"], dpi=150); plt.close(fig)

    if ev.empty:
        return out

    fig, axes = plt.subplots(len(gens), 2, figsize=(9, 4 * len(gens)), squeeze=False)
    for i, g in enumerate(gens):
        s = ev[ev["generation"] == g]
        for j, (xc, yc, lab) in enumerate((("beta_nocov", "beta_base", "no covariates -> base"),
                                           ("beta_base", "beta_full", "base -> full"))):
            ax = axes[i, j]
            if len(s):
                for st in EVALUABLE:
                    q = s[s["status"] == st]
                    ax.scatter(q[xc], q[yc], s=22, color=STATUS_COLORS[st], label=st, alpha=0.85)
                lo = float(min(s[xc].min(), s[yc].min())); hi = float(max(s[xc].max(), s[yc].max()))
                ax.plot([lo, hi], [lo, hi], color="k", lw=0.8, ls="--")
                ax.legend(frameon=False, fontsize=7)
            else:
                ax.text(0.5, 0.5, "no evaluable pair", ha="center", va="center", transform=ax.transAxes)
            ax.set_xlabel(f"G:E beta ({xc[5:]})"); ax.set_ylabel(f"G:E beta ({yc[5:]})")
            ax.set_title(f"generation {g}: {lab}", fontsize=10)
    fig.tight_layout()
    out["beta_scatter"] = fig_dir / "beta_base_vs_full.png"; fig.savefig(out["beta_scatter"], dpi=150); plt.close(fig)

    fig, ax = plt.subplots(figsize=(6, 4))
    for g in gens:
        c = fit.loc[fit["generation"] == g, "cond_full"].replace(np.inf, np.nan).dropna()
        if len(c):
            ax.hist(np.log10(c), bins=20, histtype="step", lw=1.6, label=f"gen {g}")
    ax.axvline(np.log10(COND_MAX), color="r", ls="--", lw=1)
    ax.set_xlabel("log10 condition number (full model)"); ax.set_ylabel("no. of variants")
    ax.legend(frameon=False); fig.tight_layout()
    out["cond"] = fig_dir / "condition_number.png"; fig.savefig(out["cond"], dpi=150); plt.close(fig)

    for g in gens:
        s = ev[ev["generation"] == g].copy()
        if s.empty:
            continue
        s["_k"] = s["delta_beta_full"].abs().fillna(0)
        s = s.sort_values("_k", ascending=False).head(40).iloc[::-1]
        fig, ax = plt.subplots(figsize=(8, max(3, 0.32 * len(s) + 1.5)))
        y = np.arange(len(s))
        for off, tag, mk, c in ((-0.22, "nocov", "s", "#8d99ae"), (0, "base", "o", "#264653"), (0.22, "full", "^", "#e76f51")):
            ax.errorbar(s[f"beta_{tag}"], y + off, xerr=1.96 * s[f"se_{tag}"], fmt=mk, ms=4, lw=0.8, color=c, label=tag)
        ax.axvline(0, color="k", lw=0.6)
        ax.set_yticks(y); ax.set_yticklabels([f"{v} | {e}" for v, e in zip(s["variant"], s["exposure"])], fontsize=6)
        ax.set_xlabel("G:E beta (95% CI, HC3)"); ax.set_title(f"generation {g} - largest change base -> full", fontsize=10)
        ax.legend(frameon=False, fontsize=7); fig.tight_layout()
        out[f"forest_g{g}"] = fig_dir / f"beta_ladder_gen{g}.png"; fig.savefig(out[f"forest_g{g}"], dpi=150); plt.close(fig)

    if has_perm and "perm_valid_frac" in fit.columns and fit["perm_valid_frac"].notna().any():
        fig, axes = plt.subplots(1, 2, figsize=(9, 4))
        axes[0].hist(fit["perm_valid_frac"].dropna(), bins=20, color="#264653")
        axes[0].axvline(PERM_VALID_WARN, color="r", ls="--", lw=1)
        axes[0].set_xlabel("fraction of valid permutations"); axes[0].set_ylabel("no. of variants")
        if "perm_rankdef_frac" in fit.columns:
            axes[1].hist(fit["perm_rankdef_frac"].dropna(), bins=20, color="#e76f51")
        axes[1].axvline(PERM_RANKDEF_WARN, color="r", ls="--", lw=1)
        axes[1].set_xlabel("fraction of rank-deficient permutations")
        fig.tight_layout()
        out["perm"] = fig_dir / "permutation_health.png"; fig.savefig(out["perm"], dpi=150); plt.close(fig)
    return out


# ============================================================================
# Word report
# ============================================================================
def _fmt(v):
    if isinstance(v, str):
        return v
    if v is None or (isinstance(v, (float, np.floating)) and np.isnan(v)):
        return "-"
    if isinstance(v, (bool, np.bool_)):
        return "yes" if v else "no"
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    if f.is_integer() and abs(f) < 1e9:
        return str(int(f))
    return f"{f:.1e}" if (abs(f) < 1e-3 or abs(f) >= 1e4) else f"{f:.3g}"


def build_docx(res, summ, sens, meta, wide, figs, path: Path, info: dict) -> None:
    from docx import Document
    from docx.shared import Pt
    from gene_environment.report.word_utils import (
        add_figure_to_doc, repeat_header_row, set_cell_bg, set_cell_text, set_landscape, set_table_borders,
    )

    doc = Document()
    set_landscape(doc)
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Keller sensitivity and fit diagnostics", level=0)
    doc.add_paragraph(
        f"Generated on {info['timestamp']}. Generations: {', '.join(map(str, info['generations']))}. "
        f"Variant x exposure pairs: {info['n_pairs']}. Bootstrap B = {info['boot_B']}, "
        f"permutation diagnostics: {info['perm_B'] if info['perm_B'] else 'not run'}."
    )
    doc.add_paragraph("Genetic sources: " + "; ".join(f"gen {g} -> {p}" for g, p in info["sources"].items()))

    def table(df_, columns, font_pt=8):
        t = doc.add_table(rows=1, cols=len(columns))
        set_table_borders(t)
        for i, (_, h) in enumerate(columns):
            set_cell_text(t.rows[0].cells[i], h, bold=True); set_cell_bg(t.rows[0].cells[i])
        repeat_header_row(t)
        for _, r in df_.iterrows():
            cells = t.add_row().cells
            for i, (k, _) in enumerate(columns):
                set_cell_text(cells[i], _fmt(r[k]) if k in r.index else "-")
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(font_pt)
        return t

    doc.add_heading("Models and decision rules", level=1)
    for t in (
        "Models on the same matched sample: base (pipeline: G*E + C); full (G*E + C + C:E + C:G), the main complete model.",
        f"Change: |beta_full - beta_base| / |beta_base| >= {T['CHANGE_MAIN']:.0%} and the {T['BOOT_CI']:.0%} bootstrap "
        "CI of the difference excludes 0. Sign inversion: same CI condition. Sensitivity at "
        + ", ".join(f"{x:.0%}" for x in T["CHANGE_SENSITIVITY"]) + ".",
        f"Validity: base and full fits well posed (full rank, condition < {COND_MAX:.0e}, finite HC3 SE) and SMD <= "
        f"{info['pipeline']['max_smd']}. Otherwise NOT_EVALUABLE: not evidence against the variant.",
        f"Cohort heterogeneity: Cochran Q p < {T['HET_ALPHA']}: flagged, not an exclusion. Significance with C:E / "
        "C:G is only reported.",
    ):
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("Thresholds", level=1)
    table(pd.DataFrame([{"k": k, "v": str(v)} for k, v in info["thresholds"].items()]), [("k", "threshold"), ("v", "value")])

    doc.add_heading("Summary by generation and exposure", level=1)
    cols = [("generation", "gen"), ("exposure", "exposure"), ("n_cohort", "N cohort"), ("n_pairs", "pairs"),
            ("n_stable", "STABLE"), ("n_changed", "CHANGED"), ("n_no_ci", "NO CI"), ("n_not_evaluable", "NOT EVAL."),
            ("n_skip", "SKIP"), ("n_fail", "FAIL"), ("frac_concordant", "concordant"),
            ("median_cond_full", "median cond."), ("median_abs_shift_full", "median |shift|"),
            ("median_se_ratio_full_base", "SE full/base"), ("n_sig_base", "sig base"),
            ("n_sig_full", "sig full"), ("n_sig_base_still_sig_full", "sig in both")]
    table(summ, [c for c in cols if c[0] in summ.columns])

    doc.add_heading("Sensitivity to the change threshold", level=1)
    if len(sens):
        table(sens, [("generation", "gen"), ("exposure", "exposure"), ("threshold", "threshold"),
                     ("n_evaluable", "evaluable"), ("n_change_no_ci", "change (no CI)"),
                     ("n_change_ci", "change + CI"), ("n_sign_flip_ci", "sign flip + CI")])
    else:
        doc.add_paragraph("No evaluable pair.")

    doc.add_heading("Meta-analysis across cohorts (full model)", level=1)
    if len(meta) and (meta["k"] >= 2).any():
        doc.add_paragraph(f"Pairs with >= 2 cohorts: {int((meta['k'] >= 2).sum())}; heterogeneity flagged "
                          f"(Q p < {T['HET_ALPHA']}): {int(meta['flag_het'].sum())}. Full table in all_variants_table.csv.")
    else:
        doc.add_paragraph("Fewer than two cohorts evaluable for every pair: no heterogeneity test. See meta_analysis.csv.")

    doc.add_heading("Figures", level=1)
    caps = {"status": "Status by generation and exposure.",
            "beta_scatter": "G:E beta: no covariates vs base (left), base vs full (right); evaluable pairs.",
            "cond": "Condition number of the full model; red line = threshold.",
            "perm": "Health of the diagnostic permutations."}
    for k in ("status", "beta_scatter", "cond", "perm"):
        if k in figs:
            add_figure_to_doc(doc, figs[k], caps[k], width_in=6.5)
    for k, p in figs.items():
        if k.startswith("forest_g"):
            add_figure_to_doc(doc, p, f"Model ladder, generation {k[8:]}: largest change base -> full.", width_in=7.0)

    doc.add_heading("All variants", level=1)
    gens = info["generations"]
    cols = [("variant", "variant"), ("exposure", "exposure")]
    for g in gens:
        cols += [(f"status_g{g}", f"status g{g}"), (f"rel_shift_full_g{g}", f"rel shift g{g}"),
                 (f"boot_excl0_g{g}", f"CI excl 0 g{g}")]
    cols += [("beta_fixed", "meta beta (fixed)"), ("Q_p", "Q p"), ("flag_het", "het")]
    table(wide, [c for c in cols if c[0] in wide.columns], font_pt=7)

    prob = res[res["status"].isin(["FAIL", "NOT_EVALUABLE", "SKIP"])].copy()
    doc.add_heading("Rows to review (FAIL, NOT_EVALUABLE, SKIP)", level=1)
    if prob.empty:
        doc.add_paragraph("None.")
    else:
        order = {"FAIL": 0, "NOT_EVALUABLE": 1, "SKIP": 2}
        prob = prob.sort_values(["status", "generation", "exposure"],
                                key=lambda s: s.map(order) if s.name == "status" else s)
        cut = prob.head(80)
        table(cut, [("generation", "gen"), ("exposure", "exposure"), ("variant", "variant"),
                    ("status", "status"), ("reasons", "reasons")])
        if len(prob) > len(cut):
            doc.add_paragraph(f"Showing {len(cut)} of {len(prob)} rows: full list in results_all.csv.")

    doc.add_heading("Notes", level=1)
    for t in (
        "The sample is matched with the same procedure and parameters as the pipeline "
        "(MATCH_K, matching on exposure + sex + PCs); the exposure is standardized on the generation's cohort.",
        "Bootstrap resamples the rows of the matched sample with replacement (percentile CI); the seed is fixed "
        "per variant, exposure and generation, so the run is reproducible.",
        "Thresholds are in the THRESHOLDS block at the top of keller_sensitivity.py and in run_info.json.",
    ):
        doc.add_paragraph(t, style="List Bullet")
    doc.save(str(path))


# ============================================================================
# Orchestration
# ============================================================================
def run(variants_csv: str, out_dir: str | None = None, generations=(1, 23), *, perm_B: int = 0,
        boot_B: int | None = None, sources: dict | None = None, raw_attr: str | None = None,
        allow_overlap: bool = False, dedup: bool = True, pca: dict | None = None) -> pd.DataFrame:
    from gene_environment.config import get_config
    from gene_environment.logging_utils import get_logger
    from gene_environment.vcf_pipeline.build_dataset import _build_narrow_covariates, _load_genetic_data

    log = get_logger(__name__)
    cfg = get_config()
    boot_B = T["BOOT_B"] if boot_B is None else boot_B
    if boot_B < 1:
        raise ValueError(f"bootstrap B={boot_B}: the change rule needs the bootstrap CI (default {T['BOOT_B']}). "
                         "If you call run() from a wrapper, pass boot_B by keyword.")
    sources = {**DEFAULT_SOURCES, **(sources or {})}
    missing = [g for g in generations if g not in sources]
    if missing:
        raise ValueError(f"no genetic file defined for generation(s) {missing}: use --source G=PATH")
    for g in generations:
        if not Path(sources[g]).exists():
            raise FileNotFoundError(f"generation {g}: genetic file not found: {sources[g]}")

    pca = {**DEFAULT_PCA, **(pca or {})}
    for g in generations:
        if g not in pca:
            raise ValueError(f"no PC file defined for generation {g}: use --pca G=PATH")
        if not Path(pca[g].format(generation=g)).exists():
            import glob
            raise FileNotFoundError(f"generation {g}: PC file not found: {pca[g]}. Existing: "
                                    f"{sorted(glob.glob(PCA_BASE + '/qc_output*/pca_covariates*.csv'))}")

    pairs = _read_input(variants_csv, cfg)
    out = Path(out_dir or getattr(cfg, "keller_sensitivity_dir", "./output/keller_sensitivity"))
    (out / "figures").mkdir(parents=True, exist_ok=True)
    t0 = time.time()
    log.info("Input: %d variant x exposure pairs, %d exposures, generations %s, bootstrap B=%d",
             len(pairs), pairs["exposure"].nunique(), list(generations), boot_B)

    by_path: dict[str, list[int]] = {}
    for g in generations:
        by_path.setdefault(str(sources[g]), []).append(g)

    rows = []
    ids_seen: dict[tuple, frozenset] = {}   # (exposure, generation) -> ids, for the overlap guard
    n_removed: dict[str, int] = {}
    # Overlap rule: a sample present in the file of a HIGHER generation belongs to that generation
    # and is removed from the lower one. Higher generations are processed first.
    higher_ids: set = set()
    if dedup:
        lo = min(generations)
        for g2, p2 in sources.items():
            if g2 > lo and str(p2) not in by_path and Path(p2).exists():
                log.info("Reading ids of %s (generation %s, not analysed) to remove overlaps", p2, g2)
                higher_ids |= _pipeline_ids(cfg, str(p2), raw_attr, _load_genetic_data)
    for path, gens_here in sorted(by_path.items(), key=lambda kv: -max(kv[1])):
        cfg_src, used_attr = _cfg_with_raw(cfg, path, raw_attr)
        log.info("Genetic file %s (cfg.%s) for generations %s", path, used_attr, gens_here)
        df_gen, _, mapping, _ = _load_genetic_data(cfg_src)
        orig_to_safe = {v: k for k, v in mapping.items()}
        needed = [orig_to_safe[v] for v in dict.fromkeys(pairs["variant"]) if v in orig_to_safe]
        n_missing = len(set(pairs["variant"]) - set(orig_to_safe))
        log.info("  %d samples, %d/%d requested variants present (%d absent)",
                 len(df_gen), len(needed), pairs["variant"].nunique(), n_missing)
        df_small = df_gen[["id"] + needed].copy()
        del df_gen, mapping
        gc.collect()
        if dedup:
            ids_all = df_small["id"].astype(str)
            hit = ids_all.isin(higher_ids)
            n_removed[path] = int(hit.sum())
            log.info("  generations %s: removed %d samples that belong to a higher generation; %d left",
                     gens_here, int(hit.sum()), int((~hit).sum()))
            higher_ids |= set(ids_all)          # all ids of this file, before removal
            df_small = df_small[~hit].reset_index(drop=True)

        for g in gens_here:
            for exposure, sub in pairs.groupby("exposure", sort=False):
                log.info("Generation %s, exposure %s: %d variants", g, exposure, len(sub))
                cfg_ge = dataclasses.replace(cfg_src, generation=g, exposure=exposure)
                cfg_ge, _pa = _cfg_with_attr(cfg_ge, PCA_ATTR_CANDIDATES, pca[g], "the PC file template")
                log.info("  PC file for generation %s: %s (cfg.%s)", g, pca[g], _pa)
                try:
                    cov, Ecols, Ccols = _quiet(_build_narrow_covariates, cfg_ge, df_small["id"])
                    df = pd.merge(cov, df_small, on="id", how="inner")
                except Exception as exc:
                    log.error("g%s %s: dataset could not be built: %s", g, exposure, exc)
                    rows += [{"generation": g, "exposure": exposure, "variant": v, "status": "SKIP",
                              "reasons": f"dataset could not be built: {exc}"} for v in sub["variant"]]
                    continue

                # guard: generations must not share samples
                this_ids = frozenset(df["id"].astype(str))
                for (e2, g2), ids2 in ids_seen.items():
                    if e2 == exposure and g2 != g and ids2 & this_ids:
                        msg = (f"generation {g} and {g2} share {len(ids2 & this_ids)} samples for {exposure} "
                               f"(n={len(this_ids)} and {len(ids2)})")
                        if allow_overlap:
                            log.warning(msg)
                        else:
                            raise RuntimeError(msg + ". Use --allow-overlap to ignore (or keep dedup on).")
                ids_seen[(exposure, g)] = this_ids
                log.info("  cohort g%s %s: n=%d (id examples: %s)", g, exposure, len(df), df["id"].head(3).tolist())
                nas = df[Ecols + Ccols].isna().sum()
                log.info("  complete cases (E + covariates): %d/%d; missing per column: %s",
                         int(df[Ecols + Ccols].notna().all(axis=1).sum()), len(df),
                         {k: int(v) for k, v in nas.items() if v})

                for k, lab in enumerate(sub["variant"], 1):
                    col = orig_to_safe.get(lab)
                    if col is None:
                        rows.append({"generation": g, "exposure": exposure, "variant": lab, "status": "SKIP",
                                     "reasons": f"variant not found in {Path(path).name}", "n_cohort": len(df)})
                        continue
                    try:
                        r = analyze_one(df, col, lab, exposure, g, Ecols, Ccols, cfg_ge, perm_B, boot_B)
                        r["n_cohort"] = len(df)
                        rows.append(r)
                    except Exception as exc:
                        log.exception("g%s %s %s: error", g, exposure, lab)
                        rows.append({"generation": g, "exposure": exposure, "variant": lab, "status": "FAIL",
                                     "reasons": f"error: {type(exc).__name__}: {exc}", "n_cohort": len(df)})
                    if k % 25 == 0:
                        log.info("  %d/%d", k, len(sub))
                del df, cov
                gc.collect()
        del df_small
        gc.collect()

    res = pd.DataFrame(rows)
    lead = ["generation", "exposure", "variant", "status", "reasons", "n_cohort"]
    res = res[lead + [c for c in res.columns if c not in lead]]
    max_smd_used = float(res["cfg_max_smd"].dropna().iloc[0]) if "cfg_max_smd" in res and res["cfg_max_smd"].notna().any() else float(cfg.max_smd)
    res = res.drop(columns=[c for c in ("cfg_max_smd",) if c in res.columns])
    res.to_csv(out / "results_all.csv", index=False)
    for g in generations:
        res[res["generation"] == g].to_csv(out / f"results_gen{g}.csv", index=False)

    summ = summarize(res)
    summ.to_csv(out / "summary_by_gen_exposure.csv", index=False)
    sens = sensitivity_table(res, T["CHANGE_MAIN"], T["CHANGE_SENSITIVITY"])
    sens.to_csv(out / "sensitivity_table.csv", index=False)
    meta = meta_analysis(res)
    meta.to_csv(out / "meta_analysis.csv", index=False)
    wide = all_variants_table(res, meta)
    wide.to_csv(out / "all_variants_table.csv", index=False)

    figs = make_figures(res, out / "figures", bool(perm_B))
    thresholds_saved = {k: (list(v) if isinstance(v, tuple) else v) for k, v in T.items()}
    thresholds_saved["BOOT_B"] = boot_B
    thresholds_saved["MAX_SMD (pipeline)"] = max_smd_used
    info = {"timestamp": datetime.now().strftime("%Y-%m-%d %H:%M"), "generations": list(generations),
            "sources": {g: str(sources[g]) for g in generations},
            "dedup": dedup, "n_removed_overlap": n_removed, "pca_files": {g: pca[g] for g in generations},
            "n_pairs": int(len(pairs)), "exposures": sorted(pairs["exposure"].unique().tolist()),
            "perm_B": perm_B, "boot_B": boot_B, "variants_csv": str(variants_csv),
            "thresholds": thresholds_saved,
            "pipeline": {"match_k": cfg.match_k, "min_treated": cfg.min_treated, "max_smd": max_smd_used,
                         "pca_n_components": cfg.pca_n_components, "use_pca": cfg.use_pca_covariates,
                         "random_state": cfg.random_state},
            "status_counts": res["status"].value_counts().to_dict(),
            "seconds": round(time.time() - t0, 1)}
    (out / "run_info.json").write_text(json.dumps(info, indent=2, ensure_ascii=False, default=str))
    build_docx(res, summ, sens, meta, wide, figs, out / "keller_sensitivity_report.docx", info)

    print(f"\nOutput in {out.resolve()}")
    print(summ.to_string(index=False))
    print("\nSensitivity to the change threshold:")
    print(sens.to_string(index=False))
    return res


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Keller sensitivity (full model, bootstrap, meta-analysis)")
    ap.add_argument("variants_csv", help="CSV with columns variant and exposure")
    ap.add_argument("--out-dir", default=None, help="default: cfg.keller_sensitivity_dir (./output/keller_sensitivity)")
    ap.add_argument("--generations", type=int, nargs="+", default=[1, 23],
                    help="default: 1 and 23 (23 = gen2+gen3 pooled, as GENERATION in the .env)")
    ap.add_argument("--pca", action="append", metavar="GEN=PATH", help="PC file for a generation (repeatable)")
    ap.add_argument("--source", action="append", metavar="GEN=PATH", help="genetic file for a generation (repeatable)")
    ap.add_argument("--raw-attr", default=None, help="cfg attribute holding RAW_FILE (auto-detected if omitted)")
    ap.add_argument("--allow-overlap", action="store_true", help="do not stop if two generations still share samples")
    ap.add_argument("--no-dedup", action="store_true", help="do NOT remove samples present in higher-generation files")
    ap.add_argument("--perm", type=int, default=0, metavar="B",
                    help="check the health of the first B permutations per variant (0 = skip; e.g. 500)")
    ap.add_argument("--boot", type=int, default=None, metavar="B",
                    help=f"bootstrap replicates (default {T['BOOT_B']}; use a small value only for test runs)")
    a = ap.parse_args(argv)
    sources = {}
    for s in a.source or []:
        g, p = s.split("=", 1)
        sources[int(g)] = p
    pca = {}
    for sp in a.pca or []:
        g, pth = sp.split("=", 1)
        pca[int(g)] = pth
    run(a.variants_csv, a.out_dir, tuple(a.generations), perm_B=a.perm, boot_B=a.boot, pca=pca,
        sources=sources, raw_attr=a.raw_attr, allow_overlap=a.allow_overlap, dedup=not a.no_dedup)
    return 0


if __name__ == "__main__":
    sys.exit(main())